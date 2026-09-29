# -*- coding: utf-8 -*-
"""Vygeneruje přílohu individuálního studijního plánu do dokumenty/ISP/."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "dokumenty" / "ISP" / "ISP-programovani-2-rocnik-1-pololeti-2026-2027.docx"

NAVY = "1F4E79"
NAVY_RGB = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x55, 0x55, 0x55)
ZEBRA = "F3F6F9"
HEAD = "1F4E79"


def set_run_font(run, name="Times New Roman", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_lang(run):
    rPr = run._element.get_or_add_rPr()
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "cs-CZ")
    lang.set(qn("w:eastAsia"), "cs-CZ")


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_font(run, **kwargs)
    set_lang(run)
    return run


def para(doc, text="", size=11, bold=False, italic=False, color=None, align="left",
         space_before=0, space_after=6, line=1.08):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if text:
        add_text(p, text, size=size, bold=bold, italic=italic, color=color)
    return p


def shade(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="BFBFBF", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def valign(cell, val="center"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = OxmlElement("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), val)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def cell_para(cell, text, size=10, bold=False, italic=False, color=None, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }[align]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    add_text(p, text, size=size, bold=bold, italic=italic, color=color)
    return p


def style_table(table, widths, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths) * 567)))
    tblW.set(qn("w:type"), "dxa")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 567)))
        grid.append(gc)

    for r_i, row in enumerate(table.rows):
        prevent_row_split(row)
        if header and r_i == 0:
            trPr = row._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            trPr.append(tbl_header)
        for c_i, cell in enumerate(row.cells):
            cell.width = Cm(widths[c_i])
            set_cell_border(cell, "8FA4B8" if header and r_i == 0 else "D0D7DE")
            set_cell_margins(cell)
            valign(cell, "center")
            if header and r_i == 0:
                shade(cell, HEAD)
            elif r_i % 2 == 0:
                shade(cell, ZEBRA)


def heading(doc, text):
    p = para(doc, "", size=13, bold=True, color=NAVY_RGB, space_before=12, space_after=4)
    add_text(p, text, size=13, bold=True, color=NAVY_RGB)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), NAVY)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    add_text(p, "•  ", size=11)
    add_text(p, text, size=11)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    add_text(hp, "Příloha k individuálnímu studijnímu plánu", size=9, italic=True, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(2)
    add_text(
        fp,
        "Programování  ·  2. ročník  ·  1. pololetí 2026/2027  ·  strana ",
        size=9,
        color=MUTED,
    )
    run = fp.add_run()
    set_run_font(run, size=9, color=MUTED)
    set_lang(run)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)

    para(doc, "INDIVIDUÁLNÍ STUDIJNÍ PLÁN", size=16, bold=True, color=NAVY_RGB,
         align="center", space_before=0, space_after=0)
    para(doc, "přehled učiva a podmínky klasifikace", size=12, italic=True, color=MUTED,
         align="center", space_before=0, space_after=2)
    para(doc, "Předmět Programování  ·  2. ročník  ·  1. pololetí školního roku 2026/2027",
         size=11, align="center", space_before=0, space_after=8)

    id_rows = [
        ("Žák / žákyně", "[jméno a příjmení]"),
        ("Třída", "[třída]"),
        ("Škola", "[název školy]"),
        ("Předmět", "Programování (PRG)"),
        ("Vyučující", "[jméno a příjmení vyučujícího]"),
        ("Školní rok", "2026/2027, 1. pololetí"),
        ("Období výuky", "1. září 2026 – 28. ledna 2027"),
    ]
    id_table = doc.add_table(rows=len(id_rows), cols=2)
    style_table(id_table, [4.6, 12.8], header=False)
    for i, (label, value) in enumerate(id_rows):
        shade(id_table.rows[i].cells[0], "E7EEF5")
        cell_para(id_table.rows[i].cells[0], label, size=10, bold=True, color=NAVY_RGB)
        cell_para(id_table.rows[i].cells[1], value, size=10)

    para(
        doc,
        "Tato příloha se vkládá do individuálního studijního plánu. Podepisují ji žák, zákonný zástupce a vyučující předmětu. Údaje v hranatých závorkách doplní vyučující.",
        size=10, italic=True, color=MUTED, align="justify", space_before=6, space_after=2,
    )

    heading(doc, "1.  Rozsah 1. pololetí")
    para(
        doc,
        "Ročník má 170 hodin (34 týdnů po 5 hodinách). V 1. pololetí se probírá 65 hodin, lekce 01 až 11. Jsou to dva celky — algoritmizace a objektové programování včetně pololetního projektu. SQL, tedy relační databáze, tabulky a dotazy, se probírá ve 2. pololetí.",
        align="justify", space_after=6,
    )

    heading(doc, "2.  Přehled témat")
    para(
        doc,
        "Přehled učiva 1. pololetí. Úlohy se píší v Pythonu 3.",
        align="justify", space_after=6,
    )

    topics = [
        ("A. Algoritmizace — 20 hodin", None, None, True),
        ("01", "Rekurze",
         "Vysvětlí rozdíl mezi rekurzí a cyklem. Napíše funkci s podmínkou zastavení a s rekurzivním voláním, která vrátí výsledek."),
        ("02", "Vyhledávání a řazení",
         "Najde prvek lineárním i binárním hledáním. Seřadí seznam vlastním algoritmem a porovná, jak počet kroků roste s počtem položek."),
        ("B. Objektové programování — 45 hodin", None, None, True),
        ("03", "Třídy, objekty a atributy",
         "Vysvětlí rozdíl mezi třídou a objektem. Vytvoří dva objekty téže třídy a nastaví jim atributy."),
        ("04", "Konstruktor",
         "Napíše konstruktor __init__, vysvětlí, kdy se volá, a vytvoří objekt předáním údajů."),
        ("05", "Metody a self",
         "Odliší metodu od funkce. Napíše metodu, která čte nebo mění údaje svého objektu, a zavolá metodu jiného objektu."),
        ("06", "Speciální metody a vlastnosti",
         "Napíše __str__ tak, aby se objekt vypsal čitelně. Nastaví a přečte stav objektu přes metodu."),
        ("07", "Dědičnost",
         "Zapíše vztah „je“ (například pes je zvíře). Potomek použije údaje a metody rodiče a doplní vlastní konstruktor přes super()."),
        ("08", "Polymorfismus",
         "U potomka přepíše metodu rodiče. Stejné volání metody dá u různých potomků různý výsledek. Projde seznam objektů různých typů jedním cyklem."),
        ("09", "Statické metody a proměnné",
         "Odliší údaj jednoho objektu od údaje společného pro celou třídu. Napíše statickou metodu bez self."),
        ("10", "Vlastní výjimky a iterace",
         "Napíše vlastní výjimku, vyvolá ji a odchytí. Upraví objekt tak, aby šel projít cyklem for."),
        ("11", "Návrh a pololetní projekt",
         "Ze zadání nejdřív navrhne třídy a teprve potom píše kód. Odevzdá jednu konzolovou aplikaci, která spojí učivo lekcí 03–10. Téma schvaluje vyučující."),
    ]

    table = doc.add_table(rows=1 + len(topics), cols=3)
    style_table(table, [1.5, 5.2, 10.7], header=True)
    for i, h in enumerate(["Č.", "Téma", "Co se v tématu probírá"]):
        cell_para(table.rows[0].cells[i], h, size=10, bold=True, color=RGBColor(255, 255, 255), align="center")

    for r_i, item in enumerate(topics, start=1):
        if len(item) > 3 and item[3]:
            merged = table.rows[r_i].cells[0].merge(table.rows[r_i].cells[2])
            shade(merged, "E7EEF5")
            set_cell_border(merged, "8FA4B8")
            cell_para(merged, item[0], size=10, bold=True, color=NAVY_RGB)
        else:
            cell_para(table.rows[r_i].cells[0], item[0], size=10, bold=True, align="center", color=NAVY_RGB)
            cell_para(table.rows[r_i].cells[1], item[1], size=10, bold=True)
            cell_para(table.rows[r_i].cells[2], item[2], size=10)

    heading(doc, "3.  Úkoly a klasifikace")
    bullet(doc, "Žák plní úkoly předmětu. Průběžné úkoly odevzdává v AMOS ve stejných termínech jako ostatní žáci.")
    bullet(doc, "Známkované úkoly se vypracovávají v hodině a termín je vždy hlášen předem. Žák se jich proto může zúčastnit.")
    bullet(doc, "Pololetní projekt z lekce 11 je známkovaný. Žák navrhne téma, vyučující ho schválí a žák projekt odevzdá v zadaném termínu.")
    bullet(doc, "Klasifikace za 1. pololetí vychází ze známek za známkované úkoly a ze známky za pololetní projekt, podle klasifikačního řádu školy.")
    bullet(doc, "Když z průběžné práce nevznikne dostatek známek pro uzavření klasifikace, vykoná žák v lednu individuální přezkoušení. Jinak se přezkoušení nekoná.")

    heading(doc, "4.  Individuální přezkoušení v lednu")
    para(
        doc,
        "Přezkoušení je náhradní způsob, jak doplnit chybějící známky. Vysvědčení za 1. pololetí se vydává ve čtvrtek 28. ledna 2027, proto jsou oba termíny před tímto dnem. Na státní svátek ani na prázdniny nepřipadají.",
        align="justify", space_after=6,
    )

    exams = [
        ("Termín", "Datum", "Poznámka"),
        ("Přezkoušení", "úterý 19. ledna 2027",
         "Koná se, když k tomuto datu chybí dostatek známek pro klasifikaci."),
        ("Náhradní termín", "pátek 22. ledna 2027",
         "Při omluvené neúčasti na přezkoušení. Pozdější termín už nejde stihnout před vysvědčením."),
    ]
    ex = doc.add_table(rows=len(exams), cols=3)
    style_table(ex, [4.2, 4.6, 8.6], header=True)
    for r_i, row in enumerate(exams):
        for c_i, text in enumerate(row):
            if r_i == 0:
                cell_para(ex.rows[r_i].cells[c_i], text, size=10, bold=True,
                          color=RGBColor(255, 255, 255), align="center")
            else:
                cell_para(
                    ex.rows[r_i].cells[c_i],
                    text,
                    size=10,
                    bold=(c_i == 0),
                    color=NAVY_RGB if c_i == 0 else None,
                )
                if r_i == 2:
                    shade(ex.rows[r_i].cells[c_i], "FBF9F3")

    para(doc, "Rozsah a průběh", size=11, bold=True, color=NAVY_RGB, space_before=8, space_after=2)
    bullet(doc, "Přezkoušení pokryje ta témata z přehledu výše, ze kterých žák nemá dostatek známek. Když podklady chybí z celého pololetí, zkouší se lekce 01 až 11.")
    bullet(doc, "Probíhá ve škole u počítače. Žák napíše krátký program v Pythonu a stručně vysvětlí příslušné pojmy. Trvá přibližně 45–60 minut.")
    bullet(doc, "O tom, že se přezkoušení bude konat, a o učebně informuje vyučující žáka a zákonného zástupce e-mailem nejpozději 7 dní předem.")
    bullet(doc, "Známka z přezkoušení doplní chybějící podklady. Známky za známkované úkoly a za pololetní projekt zůstávají v platnosti.")
    bullet(doc, "Výsledná známka za 1. pololetí se stanoví podle klasifikačního řádu školy.")

    doc.add_page_break()
    heading(doc, "5.  Záznam o klasifikaci")
    para(
        doc,
        "Vyučující zapíše známku za pololetní projekt. Řádek přezkoušení vyplní jen tehdy, když se přezkoušení konalo. Známky za úkoly v hodině jsou v evidenci předmětu.",
        align="justify", space_after=6,
    )
    grades = [
        ("Úkon", "Datum", "Známka", "Poznámka"),
        ("Pololetní projekt", "", "", "Známka za odevzdaný projekt."),
        ("Individuální přezkoušení", "", "", "Jen při nedostatku známek."),
        ("Výsledná známka za 1. pololetí", "", "", ""),
    ]
    gr = doc.add_table(rows=len(grades), cols=4)
    style_table(gr, [6.2, 3.4, 2.2, 5.6], header=True)
    for r_i, row in enumerate(grades):
        for c_i, text in enumerate(row):
            if r_i == 0:
                cell_para(gr.rows[r_i].cells[c_i], text, size=10, bold=True,
                          color=RGBColor(255, 255, 255), align="center")
            else:
                cell_para(
                    gr.rows[r_i].cells[c_i],
                    text,
                    size=10,
                    bold=(c_i == 0 or r_i == len(grades) - 1),
                    color=NAVY_RGB if c_i == 0 else None,
                )
                if r_i == len(grades) - 1:
                    shade(gr.rows[r_i].cells[c_i], "E7EEF5")

    heading(doc, "6.  Podpisy")
    para(
        doc,
        "Svým podpisem potvrzujeme, že jsme se seznámili s přehledem témat 1. pololetí, s možností účasti na předem hlášených známkovaných úkolech, se známkou za pololetní projekt a s podmínkami lednového přezkoušení.",
        align="justify", space_after=8,
    )

    sigs = [
        ("Vyučující předmětu", "[jméno a příjmení]"),
        ("Žák / žákyně", "[jméno a příjmení]"),
        ("Zákonný zástupce", "[jméno a příjmení]"),
    ]
    sig = doc.add_table(rows=4, cols=3)
    style_table(sig, [5.8, 5.8, 5.8], header=False)
    for i, (role, name) in enumerate(sigs):
        shade(sig.rows[0].cells[i], "E7EEF5")
        cell_para(sig.rows[0].cells[i], role, size=10, bold=True, color=NAVY_RGB, align="center")
        cell_para(sig.rows[1].cells[i], name, size=10, align="center", color=MUTED)
        cell_para(sig.rows[2].cells[i], " ", size=11)
        cell_para(sig.rows[3].cells[i], "podpis", size=9, italic=True, color=MUTED, align="center")
        set_cell_border(sig.rows[3].cells[i], "FFFFFF")
        tc = sig.rows[3].cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        top = tcPr.find(qn("w:tcBorders")).find(qn("w:top"))
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:color"), "1F4E79")

    tr = sig.rows[2]._tr
    trPr = tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), "700")
    tr_height.set(qn("w:hRule"), "atLeast")
    trPr.append(tr_height)

    para(doc, "Datum: 29. září 2026", size=11, space_before=10, space_after=2)
    para(
        doc,
        "Vypracováno podle osnovy předmětu Programování, 2. ročník (170 hodin, z toho 65 hodin v 1. pololetí).",
        size=9, italic=True, color=MUTED, space_before=0, space_after=0,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
